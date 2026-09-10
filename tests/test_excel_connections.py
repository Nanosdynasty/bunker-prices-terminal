from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

import app as module


def workbook_bytes(value="VLSFO"):
    stream = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws["A1"] = value
    ws["B1"] = 612.5
    wb.save(stream)
    return stream.getvalue()


def save_workbook(path: Path, value="VLSFO"):
    path.write_bytes(workbook_bytes(value))


def make_client(tmp_path, extra=None):
    root = tmp_path / "workbooks"
    root.mkdir()
    config = {
        "TESTING": True,
        "SECRET_KEY": "test-secret",
        "SESSION_FILE_DIR": str(tmp_path / "sessions"),
        "LOCAL_WORKBOOK_ROOT": str(root),
        "MAX_WORKBOOK_BYTES": 2 * 1024 * 1024,
        "MS_CLIENT_ID": "",
        "MS_CLIENT_SECRET": "",
    }
    if extra:
        config.update(extra)
    flask_app = module.create_app(config)
    return flask_app.test_client(), root


def csrf(client):
    client.get("/")
    with client.session_transaction() as sess:
        return sess["csrf_token"]


def test_local_folder_file_can_be_selected_and_loaded(tmp_path):
    client, root = make_client(tmp_path)
    save_workbook(root / "prices.xlsx", "HSFO")
    token = csrf(client)

    selected = client.post(
        "/api/select-file",
        json={"path": "prices.xlsx"},
        headers={"X-CSRF-Token": token},
    )
    assert selected.status_code == 200
    assert selected.get_json()["selection"]["sheet_names"] == ["Prices"]

    loaded = client.post(
        "/api/select-sheet",
        json={"worksheet": "Prices"},
        headers={"X-CSRF-Token": token},
    )
    selection = loaded.get_json()["selection"]
    assert loaded.status_code == 200
    assert selection["preview"]["rows"][0][0]["display"] == "HSFO"
    assert selection["raw_matrix"][0][1] == "612.5"


def test_onedrive_status_reports_unconfigured_by_default(tmp_path):
    client, _ = make_client(tmp_path)
    response = client.get("/api/onedrive/status")
    assert response.status_code == 200
    assert response.get_json() == {"configured": False, "connected": False, "user": ""}


def test_onedrive_file_can_be_selected_with_mocked_graph(tmp_path, monkeypatch):
    client, _ = make_client(
        tmp_path,
        {
            "MS_CLIENT_ID": "client-id",
            "MS_CLIENT_SECRET": "client-secret",
        },
    )
    token = csrf(client)

    monkeypatch.setattr(module, "graph_access_token", lambda: "token")
    monkeypatch.setattr(
        module,
        "graph_get_json",
        lambda endpoint, token, prefer=None: {
            "id": "item-id",
            "name": "cloud-prices.xlsx",
            "lastModifiedDateTime": "2026-09-09T09:00:00Z",
            "eTag": "etag-1",
            "size": 1234,
        },
    )
    monkeypatch.setattr(module, "graph_download_workbook", lambda drive_id, item_id, token: workbook_bytes("MGO"))

    selected = client.post(
        "/api/select-onedrive-file",
        json={"drive_id": "drive-id", "item_id": "item-id"},
        headers={"X-CSRF-Token": token},
    )
    assert selected.status_code == 200
    assert selected.get_json()["selection"]["source"] == "onedrive"

    loaded = client.post(
        "/api/select-sheet",
        json={"worksheet": "Prices"},
        headers={"X-CSRF-Token": token},
    )
    selection = loaded.get_json()["selection"]
    assert loaded.status_code == 200
    assert selection["filename"] == "cloud-prices.xlsx"
    assert selection["preview"]["rows"][0][0]["display"] == "MGO"


def test_sharing_url_is_encoded_for_graph():
    encoded = module.encode_sharing_url("https://contoso.sharepoint.com/:x:/r/sites/team/prices.xlsx?d=abc")
    assert encoded.startswith("u!")
    assert "=" not in encoded
    assert "/" not in encoded
    assert "+" not in encoded


def test_shared_link_file_can_be_selected_with_mocked_graph(tmp_path, monkeypatch):
    client, _ = make_client(
        tmp_path,
        {
            "MS_CLIENT_ID": "client-id",
            "MS_CLIENT_SECRET": "client-secret",
        },
    )
    token = csrf(client)

    monkeypatch.setattr(module, "graph_access_token", lambda: "token")

    def fake_graph_get_json(endpoint, token, prefer=None):
        assert endpoint.startswith("/shares/u!")
        assert prefer == "redeemSharingLinkIfNecessary"
        return {
            "id": "shared-item-id",
            "name": "shared-prices.xlsx",
            "lastModifiedDateTime": "2026-09-10T09:00:00Z",
            "eTag": "shared-etag-1",
            "size": 2345,
            "parentReference": {"driveId": "drive-id"},
            "file": {},
        }

    monkeypatch.setattr(module, "graph_get_json", fake_graph_get_json)
    monkeypatch.setattr(module, "graph_download_shared_workbook", lambda share_id, token: workbook_bytes("SHARED"))

    selected = client.post(
        "/api/select-shared-link",
        json={"url": "https://contoso.sharepoint.com/sites/team/Shared%20Documents/shared-prices.xlsx"},
        headers={"X-CSRF-Token": token},
    )
    selection = selected.get_json()["selection"]
    assert selected.status_code == 200
    assert selection["source"] == "shared_link"
    assert selection["filename"] == "shared-prices.xlsx"

    loaded = client.post(
        "/api/select-sheet",
        json={"worksheet": "Prices"},
        headers={"X-CSRF-Token": token},
    )
    selection = loaded.get_json()["selection"]
    assert loaded.status_code == 200
    assert selection["preview"]["rows"][0][0]["display"] == "SHARED"
