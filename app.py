import os
import secrets
import threading
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

import msal
import requests
from flask import Flask, jsonify, redirect, render_template, request, session, url_for
from flask_session import Session
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

try:
    from cachelib.file import FileSystemCache
except ImportError:
    FileSystemCache = None

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

EXCEL_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"}
MAX_ROWS = 2000
MAX_COLUMNS = 200
DEFAULT_WORKBOOK_ROOT = r"C:\Users\deepak\OneDrive\onedrivebunker"
GRAPH_ROOT = "https://graph.microsoft.com/v1.0"
GRAPH_SCOPES = ["User.Read", "Files.Read", "offline_access"]
_locks_guard = threading.Lock()
_session_locks: dict[str, threading.Lock] = {}


class AppError(Exception):
    def __init__(self, message, status=400, code="local_file_error"):
        super().__init__(message)
        self.message = message
        self.status = status
        self.code = code


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def create_app(test_config=None):
    app = Flask(__name__, static_folder='.', static_url_path='', template_folder='.')
    session_dir = Path(os.getenv("SESSION_FILE_DIR", Path(app.instance_path) / "sessions"))
    session_dir.mkdir(parents=True, exist_ok=True)
    
    # Resolve default workbook root: fallback to user home OneDrive or Downloads if DEFAULT_WORKBOOK_ROOT doesn't exist
    configured_root = os.getenv("LOCAL_WORKBOOK_ROOT", DEFAULT_WORKBOOK_ROOT)
    if not Path(configured_root).is_dir():
        user_onedrive = Path.home() / "OneDrive"
        downloads = Path.home() / "Downloads"
        if user_onedrive.is_dir():
            configured_root = str(user_onedrive)
        elif downloads.is_dir():
            configured_root = str(downloads)
        else:
            configured_root = str(Path.home())

    app.config.from_mapping(
        SECRET_KEY=os.getenv("FLASK_SECRET_KEY", secrets.token_hex(32)),
        SESSION_TYPE="cachelib" if FileSystemCache else "filesystem",
        SESSION_FILE_DIR=str(session_dir),
        SESSION_PERMANENT=False,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=False,
        MAX_WORKBOOK_BYTES=int(os.getenv("MAX_WORKBOOK_MB", "50")) * 1024 * 1024,
        LOCAL_WORKBOOK_ROOT=configured_root,
        MS_CLIENT_ID=os.getenv("MS_CLIENT_ID", ""),
        MS_CLIENT_SECRET=os.getenv("MS_CLIENT_SECRET", ""),
        MS_AUTHORITY=os.getenv("MS_AUTHORITY", "https://login.microsoftonline.com/common"),
        MS_REDIRECT_PATH=os.getenv("MS_REDIRECT_PATH", "/auth/callback"),
        TESTING=False,
    )
    if test_config:
        app.config.update(test_config)
    if FileSystemCache:
        app.config["SESSION_CACHELIB"] = FileSystemCache(app.config["SESSION_FILE_DIR"], threshold=500)
    Session(app)

    @app.before_request
    def protect_posts():
        if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
            supplied = request.headers.get("X-CSRF-Token") or request.form.get("csrf_token")
            if not supplied or not secrets.compare_digest(supplied, session.get("csrf_token", "")):
                return jsonify(error="Security check failed. Reload the page and try again.", code="csrf_failed"), 403

    @app.after_request
    def security_headers(response):
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "same-origin"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self' 'unsafe-inline' 'unsafe-eval' https: data:; "
            "script-src 'self' 'unsafe-inline' 'unsafe-eval' https:; "
            "style-src 'self' 'unsafe-inline' https:; "
            "img-src 'self' data: https:; "
            "font-src 'self' data: https:; "
            "connect-src 'self' https:; frame-ancestors 'none'"
        )
        return response

    @app.get("/")
    def index():
        session.setdefault("csrf_token", secrets.token_urlsafe(32))
        root = workbook_root()
        return render_template(
            "index.html", state=public_state(), csrf_token=session["csrf_token"],
            configured_root=str(root), root_available=root.is_dir(),
        )

    @app.get("/health")
    def health():
        return jsonify(
            status="ok",
            workbook_root_available=workbook_root().is_dir(),
            microsoft_auth_configured=auth_configured(),
        )

    @app.get("/api/state")
    def api_state():
        session.setdefault("csrf_token", secrets.token_urlsafe(32))
        res = public_state()
        res["csrf_token"] = session["csrf_token"]
        return jsonify(res)

    @app.get("/auth/connect")
    def auth_connect():
        if not auth_configured():
            raise AppError(
                "Microsoft sign-in is not configured. Set MS_CLIENT_ID, MS_CLIENT_SECRET, and the redirect URI first.",
                503,
                "microsoft_auth_not_configured",
            )
        flow = build_msal_app().initiate_auth_code_flow(
            scopes=GRAPH_SCOPES,
            redirect_uri=external_redirect_uri(),
        )
        session["auth_flow"] = flow
        session.modified = True
        return redirect(flow["auth_uri"])

    @app.get("/auth/callback")
    def auth_callback():
        flow = session.get("auth_flow")
        if not flow:
            raise AppError("Microsoft sign-in session expired. Start the connection again.", 400, "expired_oauth_session")
        try:
            msal_app = build_msal_app()
            result = msal_app.acquire_token_by_auth_code_flow(flow, dict(request.args))
        except ValueError as exc:
            raise AppError("Microsoft sign-in state validation failed. Start the connection again.", 400, "oauth_state_failed") from exc
        if "access_token" not in result:
            raise AppError(
                result.get("error_description") or "Microsoft permission was denied or sign-in failed.",
                403,
                "microsoft_permission_denied",
            )
        session["token_cache"] = msal_app.token_cache.serialize()
        session["account"] = result.get("id_token_claims", {})
        session.pop("auth_flow", None)
        session.modified = True
        return redirect(url_for("index"))

    @app.get("/api/onedrive/status")
    def onedrive_status():
        return jsonify({
            "configured": auth_configured(),
            "connected": graph_connected(),
            "user": signed_in_user_label(),
        })

    @app.get("/api/onedrive/folder")
    def onedrive_folder():
        token = graph_access_token()
        drive_id = request.args.get("drive_id", "")
        item_id = request.args.get("item_id", "root")
        if drive_id and item_id != "root":
            endpoint = f"/drives/{drive_id}/items/{item_id}/children"
        elif drive_id:
            endpoint = f"/drives/{drive_id}/root/children"
        else:
            endpoint = "/me/drive/root/children"
        data = graph_get_json(endpoint, token)
        items = []
        for item in data.get("value", []):
            is_folder = "folder" in item
            is_xlsx = item.get("name", "").lower().endswith(".xlsx")
            if is_folder or is_xlsx:
                items.append({
                    "name": item.get("name", ""),
                    "id": item.get("id", ""),
                    "drive_id": item.get("parentReference", {}).get("driveId") or drive_id,
                    "folder": is_folder,
                    "modified": item.get("lastModifiedDateTime"),
                    "size": item.get("size"),
                })
        items.sort(key=lambda item: (not item["folder"], item["name"].lower()))
        return jsonify(items=items, path=item_id)

    @app.post("/api/select-onedrive-file")
    def select_onedrive_file():
        payload = request.get_json(silent=True) or {}
        drive_id = payload.get("drive_id")
        item_id = payload.get("item_id")
        if not drive_id or not item_id:
            raise AppError("Missing OneDrive drive or file ID. Select the file from the browser list.", 400, "missing_drive_item")
        token = graph_access_token()
        metadata = graph_get_json(f"/drives/{drive_id}/items/{item_id}", token)
        if not metadata.get("name", "").lower().endswith(".xlsx"):
            raise AppError("Only standard .xlsx workbook files are supported.", 415, "unsupported_file")
        content = graph_download_workbook(drive_id, item_id, token)
        signature = graph_signature(metadata)
        session["selection"] = {
            "source": "onedrive",
            "drive_id": drive_id,
            "item_id": item_id,
            "filename": metadata.get("name", "OneDrive workbook.xlsx"),
            "file_modified": metadata.get("lastModifiedDateTime"),
            "observed_signature": signature,
            "loaded_signature": None,
            "sheet_names": workbook_sheet_names(content),
            "worksheet": None,
            "last_check": utc_now(),
            "last_load": None,
            "changed_detected": None,
            "stale": False,
            "error": None,
            "preview": None,
            "raw_matrix": None,
        }
        session.modified = True
        return jsonify(public_state())

    @app.get("/api/folder")
    def list_folder():
        relative = request.args.get("path", "")
        folder = safe_path(relative, require_exists=True)
        if not folder.is_dir():
            raise AppError("The selected path is not a folder.", 400, "invalid_folder")
        items = []
        try:
            for child in folder.iterdir():
                if child.is_symlink():
                    continue
                if child.is_dir() or (child.is_file() and child.suffix.lower() == ".xlsx"):
                    items.append({
                        "name": child.name,
                        "path": child.relative_to(workbook_root()).as_posix(),
                        "folder": child.is_dir(),
                    })
        except OSError as exc:
            raise AppError("The configured folder could not be read.", 403, "folder_unavailable") from exc
        items.sort(key=lambda item: (not item["folder"], item["name"].lower()))
        return jsonify(path=relative, items=items)

    @app.post("/api/select-file")
    def select_file():
        relative = (request.get_json(silent=True) or {}).get("path", "")
        path = validate_workbook_path(relative)
        content, stat = read_workbook(path)
        session["selection"] = {
            "relative_path": path.relative_to(workbook_root()).as_posix(),
            "filename": path.name,
            "file_modified": modified_iso(stat),
            "observed_signature": file_signature(stat),
            "loaded_signature": None,
            "sheet_names": workbook_sheet_names(content),
            "worksheet": None,
            "last_check": utc_now(),
            "last_load": None,
            "changed_detected": None,
            "stale": False,
            "error": None,
            "preview": None,
            "raw_matrix": None,
        }
        session.modified = True
        return jsonify(public_state())

    @app.post("/api/select-sheet")
    def select_sheet():
        selected = require_selection()
        sheet = (request.get_json(silent=True) or {}).get("worksheet")
        content, metadata = read_selected_workbook(selected)
        names = workbook_sheet_names(content)
        if sheet not in names:
            raise AppError("That worksheet does not exist in the workbook.", 409, "missing_worksheet")
        
        raw_rows = workbook_raw_matrix(content, sheet)
        selected.update(
            worksheet=sheet,
            preview=workbook_preview(content, sheet),
            raw_matrix=raw_rows,
            sheet_names=names,
            file_modified=metadata["modified"],
            observed_signature=metadata["signature"],
            loaded_signature=metadata["signature"],
            last_check=utc_now(),
            last_load=utc_now(),
            changed_detected=False,
            stale=False,
            error=None,
        )
        session.modified = True
        return jsonify(public_state())

    @app.post("/api/refresh")
    def refresh():
        selected = require_selection()
        if not selected.get("worksheet"):
            raise AppError("Select a worksheet before refreshing.", 400, "missing_worksheet")
        payload = request.get_json(silent=True) or {}
        force_reload = payload.get("force", True) is not False
        lock = session_lock()
        if not lock.acquire(blocking=False):
            raise AppError("A refresh is already running for this session.", 409, "refresh_in_progress")
        try:
            return jsonify(run_refresh(selected, force_reload=force_reload))
        finally:
            lock.release()

    @app.post("/api/change-file")
    def change_file():
        session.pop("selection", None)
        return jsonify(public_state())

    @app.post("/api/disconnect")
    def disconnect():
        session.clear()
        session["csrf_token"] = secrets.token_urlsafe(32)
        return jsonify(public_state(), csrf_token=session["csrf_token"])

    @app.errorhandler(AppError)
    def handle_app_error(error):
        return jsonify(error=error.message, code=error.code, state=public_state()), error.status

    def run_refresh(selected, force_reload=False):
        try:
            metadata = selected_metadata(selected)
            signature = metadata["signature"]
            selected["last_check"] = utc_now()
            selected["file_modified"] = metadata["modified"]
            selected["observed_signature"] = signature
            changed = signature != selected.get("loaded_signature")
            selected["changed_detected"] = changed
            if changed or force_reload:
                content, metadata = read_selected_workbook(selected)
                names = workbook_sheet_names(content)
                if selected["worksheet"] not in names:
                    raise AppError("The selected worksheet is missing from the changed workbook.", 409, "missing_worksheet")
                selected["preview"] = workbook_preview(content, selected["worksheet"])
                selected["raw_matrix"] = workbook_raw_matrix(content, selected["worksheet"])
                selected["sheet_names"] = names
                selected["loaded_signature"] = metadata["signature"]
                selected["last_load"] = utc_now()
                selected["file_modified"] = metadata["modified"]
            selected["stale"] = False
            selected["error"] = None
            session.modified = True
            return public_state()
        except (AppError, OSError) as exc:
            error = exc if isinstance(exc, AppError) else AppError(
                "The local workbook could not be read. It may be missing or temporarily locked by Excel.",
                503, "file_unavailable",
            )
            selected["stale"] = selected.get("preview") is not None
            selected["error"] = error.message
            session.modified = True
            raise error

    return app


def workbook_root():
    from flask import current_app
    return Path(current_app.config["LOCAL_WORKBOOK_ROOT"]).expanduser().resolve()


def safe_path(relative, require_exists=False):
    root = workbook_root()
    candidate = (root / Path(relative or ".")).resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise AppError("That path is outside the configured workbook folder.", 403, "path_not_allowed") from exc
    if require_exists and not candidate.exists():
        raise AppError("The selected local file or folder no longer exists.", 404, "missing_file")
    return candidate


def validate_workbook_path(relative):
    path = safe_path(relative, require_exists=True)
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise AppError("Only standard .xlsx workbook files are supported.", 415, "unsupported_file")
    return path


def read_workbook(path):
    from flask import current_app
    try:
        stat = path.stat()
        if stat.st_size > int(current_app.config["MAX_WORKBOOK_BYTES"]):
            raise AppError("The workbook is larger than this prototype allows.", 413, "file_too_large")
        return path.read_bytes(), stat
    except AppError:
        raise
    except OSError as exc:
        raise AppError(
            "The local workbook could not be read. It may be missing or temporarily locked by Excel.",
            503, "file_unavailable",
        ) from exc


def read_selected_workbook(selected):
    if selected.get("source") == "onedrive":
        token = graph_access_token()
        metadata = graph_get_json(f"/drives/{selected['drive_id']}/items/{selected['item_id']}", token)
        content = graph_download_workbook(selected["drive_id"], selected["item_id"], token)
        return content, {"modified": metadata.get("lastModifiedDateTime"), "signature": graph_signature(metadata)}
    path = validate_workbook_path(selected["relative_path"])
    content, stat = read_workbook(path)
    return content, {"modified": modified_iso(stat), "signature": file_signature(stat)}


def selected_metadata(selected):
    if selected.get("source") == "onedrive":
        token = graph_access_token()
        metadata = graph_get_json(f"/drives/{selected['drive_id']}/items/{selected['item_id']}", token)
        return {"modified": metadata.get("lastModifiedDateTime"), "signature": graph_signature(metadata)}
    path = validate_workbook_path(selected["relative_path"])
    stat = path.stat()
    return {"modified": modified_iso(stat), "signature": file_signature(stat)}


def graph_signature(metadata):
    return ":".join(str(metadata.get(key, "")) for key in ("eTag", "lastModifiedDateTime", "size"))


def build_msal_app():
    from flask import current_app
    cache = msal.SerializableTokenCache()
    if session.get("token_cache"):
        cache.deserialize(session["token_cache"])
    return msal.ConfidentialClientApplication(
        current_app.config["MS_CLIENT_ID"],
        authority=current_app.config["MS_AUTHORITY"],
        client_credential=current_app.config["MS_CLIENT_SECRET"],
        token_cache=cache,
    )


def auth_configured():
    from flask import current_app
    return bool(current_app.config["MS_CLIENT_ID"] and current_app.config["MS_CLIENT_SECRET"])


def external_redirect_uri():
    return os.getenv("MS_REDIRECT_URI") or url_for("auth_callback", _external=True)


def graph_connected():
    return bool(session.get("token_cache"))


def signed_in_user_label():
    account = session.get("account") or {}
    return account.get("preferred_username") or account.get("name") or ""


def graph_access_token():
    if not auth_configured():
        raise AppError("Microsoft sign-in is not configured on this server.", 503, "microsoft_auth_not_configured")
    app = build_msal_app()
    accounts = app.get_accounts()
    result = app.acquire_token_silent(GRAPH_SCOPES, account=accounts[0] if accounts else None)
    if not result:
        raise AppError("Microsoft session expired. Connect OneDrive again.", 401, "microsoft_session_expired")
    if "access_token" not in result:
        raise AppError(result.get("error_description") or "Could not renew Microsoft access.", 401, "microsoft_token_failed")
    session["token_cache"] = app.token_cache.serialize()
    session.modified = True
    return result["access_token"]


def graph_headers(token):
    return {"Authorization": f"Bearer {token}"}


def graph_get_json(endpoint, token):
    response = requests.get(f"{GRAPH_ROOT}{endpoint}", headers=graph_headers(token), timeout=20)
    if response.status_code == 401:
        raise AppError("Microsoft session expired. Connect OneDrive again.", 401, "microsoft_session_expired")
    if response.status_code == 403:
        raise AppError("Microsoft denied access to that OneDrive item.", 403, "microsoft_permission_denied")
    if response.status_code == 404:
        raise AppError("The selected OneDrive file or folder no longer exists.", 404, "missing_file")
    if response.status_code == 429:
        raise AppError("Microsoft Graph is throttling requests. Wait a minute and refresh again.", 429, "graph_throttled")
    if not response.ok:
        raise AppError("Microsoft Graph request failed.", 502, "graph_request_failed")
    return response.json()


def graph_download_workbook(drive_id, item_id, token):
    from flask import current_app
    response = requests.get(
        f"{GRAPH_ROOT}/drives/{drive_id}/items/{item_id}/content",
        headers=graph_headers(token),
        timeout=60,
        allow_redirects=True,
    )
    if response.status_code == 401:
        raise AppError("Microsoft session expired. Connect OneDrive again.", 401, "microsoft_session_expired")
    if response.status_code == 403:
        raise AppError("Microsoft denied access to download that workbook.", 403, "microsoft_permission_denied")
    if response.status_code == 404:
        raise AppError("The selected OneDrive workbook no longer exists.", 404, "missing_file")
    if response.status_code == 429:
        raise AppError("Microsoft Graph is throttling downloads. Wait a minute and refresh again.", 429, "graph_throttled")
    if not response.ok:
        raise AppError("OneDrive workbook download failed.", 502, "download_failed")
    if len(response.content) > int(current_app.config["MAX_WORKBOOK_BYTES"]):
        raise AppError("The workbook is larger than this prototype allows.", 413, "file_too_large")
    return response.content


def file_signature(stat):
    return f"{stat.st_mtime_ns}:{stat.st_size}"


def modified_iso(stat):
    return datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(timespec="seconds")


def workbook_sheet_names(content):
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        names = workbook.sheetnames
        workbook.close()
        return names
    except Exception as exc:
        raise AppError("The file is not a readable .xlsx workbook.", 422, "invalid_workbook") from exc


def workbook_raw_matrix(content, sheet_name):
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[sheet_name]
        matrix = []
        for row in ws.iter_rows(values_only=True):
            r = []
            for cell in row:
                if isinstance(cell, (datetime, date)):
                    r.append(cell.isoformat())
                elif cell is None:
                    r.append("")
                else:
                    r.append(str(cell))
            matrix.append(r)
        wb.close()
        return matrix
    except Exception:
        return []


def workbook_preview(content, sheet_name):
    try:
        values_wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        formulas_wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
        if sheet_name not in values_wb.sheetnames:
            raise AppError("The selected worksheet does not exist.", 409, "missing_worksheet")
        values_ws, formulas_ws = values_wb[sheet_name], formulas_wb[sheet_name]
        row_count = min(max(values_ws.max_row, formulas_ws.max_row, 1), MAX_ROWS)
        natural_columns = max(values_ws.max_column, formulas_ws.max_column, 1)
        column_count = min(natural_columns, MAX_COLUMNS)
        rows = []
        value_rows = values_ws.iter_rows(
            min_row=1, max_row=row_count, min_col=1, max_col=column_count, values_only=False
        )
        formula_rows = formulas_ws.iter_rows(
            min_row=1, max_row=row_count, min_col=1, max_col=column_count, values_only=False
        )
        for value_row, formula_row in zip(value_rows, formula_rows):
            row = []
            for value_cell, formula_cell in zip(value_row, formula_row):
                value = value_cell.value
                missing_cache = formula_cell.data_type == "f" and value is None
                is_error = value_cell.data_type == "e" or (isinstance(value, str) and value in EXCEL_ERRORS)
                if missing_cache:
                    display, kind = "[missing cached result]", "missing-cache"
                elif is_error:
                    display, kind = f"[Excel error: {value}]", "error"
                else:
                    display, kind = display_value(value), "value"
                row.append({"display": display, "kind": kind})
            rows.append(row)
        values_wb.close()
        formulas_wb.close()
        return {
            "columns": [get_column_letter(i) for i in range(1, column_count + 1)],
            "rows": rows,
            "row_count": row_count,
            "column_count": column_count,
            "columns_truncated": natural_columns > MAX_COLUMNS,
        }
    except AppError:
        raise
    except Exception as exc:
        raise AppError("The worksheet could not be read. The previous preview is unchanged.", 422, "invalid_workbook") from exc


def display_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def require_selection():
    selected = session.get("selection")
    if not selected:
        raise AppError("Choose a local workbook first.", 400, "missing_selection")
    return selected


def session_lock():
    key = str(getattr(session, "sid", "anonymous"))
    with _locks_guard:
        return _session_locks.setdefault(key, threading.Lock())


def public_state():
    root = workbook_root()
    return {
        "connected": True,
        "configured_root": str(root),
        "root_available": root.is_dir(),
        "onedrive_connected": graph_connected(),
        "onedrive_user": signed_in_user_label(),
        "microsoft_auth_configured": auth_configured(),
        "selection": session.get("selection"),
    }


app = create_app()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5050")), debug=os.getenv("FLASK_DEBUG") == "1")
